"""Ingesta del Excel resumen de nóminas devuelto por la gestoría.

Formato esperado (una fila por empleado, columnas por concepto):

  codigo_empleado | nombre | salario_fijo | salario_variable | ss_empresa |
  ss_trabajador | irpf | liquido
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from ..enums import PayrollConcept, ValidationStatus
from ..models import Employee, PayrollImport, PayrollLine
from .excel_commissions import _norm, _parse_amount


# columna_normalizada -> concepto
CONCEPT_COLUMNS: dict[str, PayrollConcept] = {
    "salario_fijo": PayrollConcept.SALARIO_FIJO,
    "fijo": PayrollConcept.SALARIO_FIJO,
    "sueldo_base": PayrollConcept.SALARIO_FIJO,
    "salario_variable": PayrollConcept.SALARIO_VARIABLE,
    "variable": PayrollConcept.SALARIO_VARIABLE,
    "comisiones": PayrollConcept.SALARIO_VARIABLE,
    "ss_empresa": PayrollConcept.SEGURIDAD_SOCIAL_EMPRESA,
    "seguridad_social_empresa": PayrollConcept.SEGURIDAD_SOCIAL_EMPRESA,
    "ss_trabajador": PayrollConcept.SEGURIDAD_SOCIAL_TRABAJADOR,
    "seguridad_social_trabajador": PayrollConcept.SEGURIDAD_SOCIAL_TRABAJADOR,
    "irpf": PayrollConcept.RETENCION_IRPF,
    "retencion_irpf": PayrollConcept.RETENCION_IRPF,
    "retencion": PayrollConcept.RETENCION_IRPF,
    "liquido": PayrollConcept.LIQUIDO_A_PAGAR,
    "liquido_a_pagar": PayrollConcept.LIQUIDO_A_PAGAR,
    "neto": PayrollConcept.LIQUIDO_A_PAGAR,
    "anticipos": PayrollConcept.ANTICIPOS,
}


def parse_workbook(path: str | Path) -> list[dict]:
    """Devuelve filas: {employee_code, concepts: {PayrollConcept: amount}}."""
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header = [_norm(c) for c in rows[0]]
    code_idx = next((i for i, h in enumerate(header) if h in ("codigo", "codigo_empleado", "empleado", "matricula")), None)
    if code_idx is None:
        raise ValueError(f"No se encuentra la columna de empleado. Cabeceras: {header}")

    concept_idx = {i: CONCEPT_COLUMNS[h] for i, h in enumerate(header) if h in CONCEPT_COLUMNS}
    if not concept_idx:
        raise ValueError(f"No se reconoce ninguna columna de concepto de nómina. Cabeceras: {header}")

    parsed: list[dict] = []
    for raw in rows[1:]:
        if raw is None or all(c is None for c in raw):
            continue
        code = raw[code_idx]
        if code in (None, ""):
            continue
        concepts: dict[PayrollConcept, float] = {}
        for idx, concept in concept_idx.items():
            amount = _parse_amount(raw[idx]) if idx < len(raw) else 0.0
            if amount:
                concepts[concept] = concepts.get(concept, 0.0) + amount
        parsed.append({"employee_code": str(code).strip(), "concepts": concepts})
    return parsed


def ingest(db: Session, *, period_id: int, filename: str, stored_path: str, uploaded_by: int | None) -> PayrollImport:
    rows = parse_workbook(stored_path)
    imp = PayrollImport(
        period_id=period_id,
        filename=filename,
        stored_path=str(stored_path),
        uploaded_by=uploaded_by,
        row_count=len(rows),
    )
    db.add(imp)
    db.flush()

    employees = {e.employee_code: e for e in db.query(Employee).all()}
    for r in rows:
        emp = employees.get(r["employee_code"])
        for concept, amount in r["concepts"].items():
            db.add(
                PayrollLine(
                    import_id=imp.id,
                    employee_code=r["employee_code"],
                    employee_id=emp.id if emp else None,
                    concept=concept,
                    amount=amount,
                    validation_status=ValidationStatus.PENDING,
                )
            )
    return imp
