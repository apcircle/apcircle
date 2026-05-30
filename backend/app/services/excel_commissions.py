"""Ingesta del Excel de comisiones extraído de los sistemas internos.

Mapeo de columnas configurable y tolerante a variaciones de cabecera.
Formato esperado (cabeceras flexibles, no sensibles a may/min ni acentos):

    codigo_empleado | nombre | departamento | concepto | importe | observaciones
"""
from __future__ import annotations

import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from ..models import CommissionImport, CommissionLine, Employee
from ..enums import ValidationStatus


# alias_de_cabecera -> campo canónico
COLUMN_ALIASES: dict[str, str] = {
    "codigo": "employee_code",
    "codigo_empleado": "employee_code",
    "cod_empleado": "employee_code",
    "empleado": "employee_code",
    "matricula": "employee_code",
    "nombre": "full_name",
    "nombre_empleado": "full_name",
    "departamento": "department",
    "depto": "department",
    "concepto": "concept",
    "tipo": "concept",
    "importe": "amount",
    "comision": "amount",
    "cantidad": "amount",
    "observaciones": "notes",
    "notas": "notes",
}


def _norm(text: str) -> str:
    text = str(text or "").strip().lower()
    text = "".join(c for c in unicodedata.normalize("NFD", text) if unicodedata.category(c) != "Mn")
    return text.replace(" ", "_")


def _parse_amount(value) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    s = str(value).strip().replace("€", "").replace(" ", "")
    # admite formato español "1.234,56" y anglosajón "1,234.56"
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return round(float(s), 2)
    except ValueError:
        return 0.0


def parse_workbook(path: str | Path) -> list[dict]:
    """Lee el Excel y devuelve filas normalizadas como dicts."""
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header = [_norm(c) for c in rows[0]]
    field_index: dict[str, int] = {}
    for idx, col in enumerate(header):
        canonical = COLUMN_ALIASES.get(col)
        if canonical and canonical not in field_index:
            field_index[canonical] = idx

    if "employee_code" not in field_index or "amount" not in field_index:
        raise ValueError(
            "El Excel de comisiones debe tener al menos columnas de empleado e importe. "
            f"Cabeceras detectadas: {header}"
        )

    parsed: list[dict] = []
    for raw in rows[1:]:
        if raw is None or all(c is None for c in raw):
            continue
        code = raw[field_index["employee_code"]]
        if code in (None, ""):
            continue
        parsed.append(
            {
                "employee_code": str(code).strip(),
                "full_name": raw[field_index["full_name"]] if "full_name" in field_index else None,
                "department": raw[field_index["department"]] if "department" in field_index else None,
                "concept": str(raw[field_index["concept"]]).strip() if "concept" in field_index else "COMISION",
                "amount": _parse_amount(raw[field_index["amount"]]),
                "notes": raw[field_index["notes"]] if "notes" in field_index else None,
            }
        )
    return parsed


def ingest(db: Session, *, period_id: int, filename: str, stored_path: str, uploaded_by: int | None) -> CommissionImport:
    """Crea CommissionImport + líneas resolviendo el empleado contra el maestro."""
    rows = parse_workbook(stored_path)
    imp = CommissionImport(
        period_id=period_id,
        filename=filename,
        stored_path=str(stored_path),
        uploaded_by=uploaded_by,
        row_count=len(rows),
    )
    db.add(imp)
    db.flush()

    # índice del maestro por código
    employees = {e.employee_code: e for e in db.query(Employee).all()}

    for r in rows:
        emp = employees.get(r["employee_code"])
        line = CommissionLine(
            import_id=imp.id,
            employee_code=r["employee_code"],
            employee_id=emp.id if emp else None,
            department=r["department"],
            concept=r["concept"],
            amount=r["amount"],
            notes=r["notes"],
            validation_status=ValidationStatus.OK if emp else ValidationStatus.UNMATCHED,
        )
        db.add(line)
    return imp
